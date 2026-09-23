#!/usr/bin/env python3
"""
Builds "2027-28 FTE Planning Database.xlsx" - the starting Google Sheet for the
2027-28 FTE Planning portal.

Every full-time slot on the 2026-27 Approved FTE Lists becomes a 2027-28 position,
updated with what ADP shows today:
  * person active in ADP          -> Returning (name + Position ID from ADP)
  * marked VACANT but still active -> Not returning
  * terminated / not found in ADP  -> Open (previous holder noted)
  * open / NEW slots               -> Open
Rows that repeat the same ADP person are skipped.

Upload the file to Google Drive, open it as a Google Sheet, and follow SETUP_2027.md.
The file contains staff names - keep it out of git (it is in .gitignore).
"""

import os
import json
import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from adp_rows import adp, SCHOOL_YEAR, ADP_HEADERS, adp_row, job_group

HERE = os.path.dirname(os.path.abspath(__file__))
OUTPUT = os.path.join(HERE, "2027-28 FTE Planning Database.xlsx")

POSITION_HEADERS = ["position_key", "school_year", "campus_code", "job_code", "assignment", "first_name", "last_name",
                    "adp_position_id", "status", "notes", "created_at", "created_by", "updated_at", "updated_by", "deleted"]
LOG_HEADERS = ["timestamp", "user", "action", "position_key", "campus_code", "job_code", "field", "old_value", "new_value", "note"]
EDITOR_HEADERS = ["email", "name", "role", "campuses", "added_at", "added_by"]
SEED_EDITORS = [
    ["adal@ssttx.org", "Ali Dal", "editor", "ALL"],
    ["hkendirci@ssttx.org", "Hasan Kendirci", "editor", "ALL"],
]
SEED_USER = "seed-import"


def main():
    with open(adp.ADP_CACHE_PATH) as f:
        cache = json.load(f)
    workers = cache["workers"]

    rows = [r for r in adp.load_fte_lists() if not r["sub_pt"]]
    report = adp.reconcile(rows, workers)
    rec_by_source = {r["source"]: r for r in report["records"] if r["source"] != "ADP"}
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    positions, log, skipped = [], [], 0
    for row in rows:
        rec = rec_by_source[f"{row['file']} (row {row['line']})"]
        st = rec["status"]
        if st == "DUPLICATE_ON_LIST":
            skipped += 1
            continue
        first = last = pid = ""
        notes = [n for n in (row["notes"],) if n]
        if st in ("MATCH", "LOCATION_MISMATCH", "TITLE_MISMATCH", "NAME_MISMATCH", "VACANT_STILL_ACTIVE"):
            w = next((x for x in workers if x["status"] == "Active"
                      and rec["adp_position_id"] in {a["position_id"] for a in x["assignments"]}), None)
            first, last = (w["given"], w["family"]) if w else (row["first"], row["last"])
            pid = rec["adp_position_id"]
            status = "Not returning" if st == "VACANT_STILL_ACTIVE" else "Returning"
            if w and st == "NAME_MISMATCH":
                notes.append(f"2026-27 list name: {row['name']}")
        else:
            status = "Open"
            if row["name"]:
                why = rec["adp_status"] or "not found in ADP"
                notes.append(f"2026-27: {row['name']} ({why})")
        key = f"P-{len(positions) + 1:05d}"
        positions.append([key, SCHOOL_YEAR, row["campus_code"], row["job_code"], row["assignment"], first, last, pid,
                          status, "; ".join(notes), now, SEED_USER, now, SEED_USER, ""])
        log.append([now, SEED_USER, "SEED_IMPORT", key, row["campus_code"], row["job_code"], "", "", status,
                    f"Seeded from 2026-27 approved list ({row['file']} row {row['line']}) + ADP pulled {cache['pulled_at']}"])

    # Job titles: everything planned or seen in ADP at campuses, minus substitutes / part-time aides
    titles = {}
    for w in workers:
        if w["status"] != "Active":
            continue
        a = adp.current_assignment(w)
        if a.get("job_code") and a["job_code"] not in adp.SUB_PT_JOB_CODES:
            titles.setdefault(a["job_code"], a.get("job_title", ""))
    for r in rows:
        titles.setdefault(r["job_code"], r["job_title"])

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    def sheet(name, headers, data, widths=None):
        ws = wb.create_sheet(name)
        ws.append(headers)
        for d in data:
            ws.append(d)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", start_color="0B2545")
            c.alignment = Alignment(vertical="center")
        ws.freeze_panes = "A2"
        for i, wdt in enumerate(widths or [], start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = wdt
        return ws

    sheet("Positions", POSITION_HEADERS, positions, [10, 9, 8, 11, 34, 16, 18, 13, 13, 40, 19, 14, 19, 14, 8])
    sheet("Change_Log", LOG_HEADERS, log, [19, 22, 16, 10, 8, 11, 14, 22, 22, 60])
    sheet("Editors", EDITOR_HEADERS, [e + [now, SEED_USER] for e in SEED_EDITORS], [26, 20, 9, 30, 19, 14])
    sheet("Campuses", ["code", "name", "region"], [[c[0], c[1], c[2]] for c in adp.CAMPUSES], [8, 36, 16])
    sheet("Job_Titles", ["job_code", "title", "group"],
          sorted([[c, t, job_group(c, t)] for c, t in titles.items()], key=lambda x: (x[2], x[0])), [12, 36, 22])
    sheet("ADP_Positions", ADP_HEADERS, [adp_row(w) for w in workers])
    sheet("Settings", ["key", "value"], [["school_year", SCHOOL_YEAR], ["adp_pulled_at", cache["pulled_at"]]], [18, 30])
    wb.save(OUTPUT)

    from collections import Counter
    c = Counter(p[8] for p in positions)
    print(f"Positions: {len(positions)}  {dict(c)}  (skipped {skipped} duplicate rows)")
    print(f"Job titles: {len(titles)} · ADP workers: {len(workers)} (pulled {cache['pulled_at']})")
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
